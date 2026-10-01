"""組會投影片（.pptx）：把排定的論文＋重點欄＋參數＋圖卡＋疑問標註做成一份可以直接改的投影片草稿。

兩種內容來源：
- 直接用重點欄（免費、立即）：把各欄位切成條列。
- AI 精簡（需設定 AI）：AI 讀重點欄、參數、標註，改寫成每頁 3–5 點的簡短要點，完整內容放在講者備忘稿。
"""
import io
import json
import re
from datetime import datetime
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import ai, config, db, figures, params
from .jobs import handler

OUT_DIR = config.DATA_DIR / "exports"
W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x1F, 0x3A, 0x4D)
INK = RGBColor(0x1E, 0x29, 0x33)
MUTED = RGBColor(0x5B, 0x6B, 0x78)
PALE = RGBColor(0xEE, 0xF3, 0xF6)
ACCENT = RGBColor(0xD9, 0x77, 0x06)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xC9, 0xD6, 0xDF)
LATIN, EA = "Calibri", "Microsoft JhengHei"
M = Inches(0.6)


# ------------------------------------------------------------------ 版面小工具
def _font(run, size, color=INK, bold=False, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = LATIN
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", EA)


def _box(slide, x, y, w, h, fill=None, radius=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    if radius:
        shp.adjustments[0] = 0.06
    return shp


def _text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, autofit=True):
    """paras：[(文字, 大小, 顏色, 粗體, 對齊, 條列)]"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    if autofit:
        bp = tf._txBody.find(qn("a:bodyPr"))
        for el in list(bp):
            if el.tag in (qn("a:spAutoFit"), qn("a:noAutofit"), qn("a:normAutofit")):
                bp.remove(el)
        bp.append(bp.makeelement(qn("a:normAutofit"), {}))
    first = True
    for item in paras:
        text, size, color, bold = item[0], item[1], item[2] if len(item) > 2 else INK, item[3] if len(item) > 3 else False
        align = item[4] if len(item) > 4 else PP_ALIGN.LEFT
        bullet = item[5] if len(item) > 5 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.28))))
            pPr.set("indent", str(-Emu(Inches(0.28))))
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bu)
            p.space_after = Pt(size * 0.55)
        else:
            p.space_after = Pt(size * 0.3)
        r = p.add_run()
        r.text = text
        _font(r, size, color, bold)
    return tb


def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _title(slide, text, sub=None, dark=False):
    _text(slide, M, Inches(0.45), W - 2 * M, Inches(0.9),
          [(text, 30, WHITE if dark else NAVY, True)] + ([(sub, 13, SOFT if dark else MUTED)] if sub else []), autofit=True)


def _footer(slide, text, dark=False):
    _text(slide, M, H - Inches(0.45), W - 2 * M, Inches(0.3), [(text, 10, SOFT if dark else MUTED, False, PP_ALIGN.RIGHT)], autofit=False)


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text[:6000]


def _picture_fit(slide, path: Path, x, y, w, h):
    """把圖片等比例放進 (x, y, w, h)，置中。"""
    with pymupdf.open(path) as im:
        iw, ih = im[0].rect.width, im[0].rect.height
    s = min(w / iw, h / ih)
    pw, ph = int(iw * s), int(ih * s)
    return slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def _page_image(con, pid: int) -> Path | None:
    """論文第一頁的預覽圖（標題頁用）。"""
    f = con.execute("SELECT id, path FROM files WHERE paper_id=? ORDER BY CASE role WHEN 'main' THEN 0 ELSE 1 END, id LIMIT 1", (pid,)).fetchone()
    if not f:
        return None
    out = OUT_DIR / "cache" / f"page1-{f['id']}.png"
    if not out.exists():
        out.parent.mkdir(parents=True, exist_ok=True)
        try:
            with pymupdf.open(config.DATA_DIR / f["path"]) as doc:
                page = doc[0]
                zoom = 900 / page.rect.width
                page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).save(out)
        except Exception:  # noqa: BLE001
            return None
    return out


def _bullets(text: str, n=5, width=90) -> list[str]:
    parts = [x.strip(" -•・") for x in re.split(r"[。；;\n]|(?<=[.!?])\s+", text or "") if x.strip(" -•・")]
    out = []
    for x in parts:
        out.append(x if len(x) <= width else x[: width - 1] + "…")
        if len(out) >= n:
            break
    return out


# ------------------------------------------------------------------ 內容
def _paper_data(con, pid: int, uid: int, fig_ids: list[int] | None) -> dict:
    p = dict(con.execute("SELECT * FROM papers WHERE id=?", (pid,)).fetchone())
    p["authors"] = json.loads(p["authors"] or "[]")
    p["keyinfo"] = json.loads(p["keyinfo"] or "{}")
    p["params"] = params.paper_params(con, pid)
    q = "SELECT * FROM figures WHERE paper_id=?"
    if fig_ids is not None:
        figs = [dict(r) for r in con.execute(q + " ORDER BY page, id", (pid,)) if r["id"] in fig_ids]
    else:
        figs = [dict(r) for r in con.execute(q + " ORDER BY kind!='figure', author_id=? DESC, page, id LIMIT 4", (pid, uid))]
    p["figures"] = figs
    p["questions"] = [dict(r) for r in con.execute(
        "SELECT a.page, a.quote, a.body, u.display_name AS who FROM annotations a LEFT JOIN users u ON u.id=a.author_id "
        "WHERE a.paper_id=? AND a.private=0 AND a.color='red' ORDER BY a.page, a.id LIMIT 6", (pid,))]
    p["highlights"] = [dict(r) for r in con.execute(
        "SELECT a.page, a.quote, a.body FROM annotations a WHERE a.paper_id=? AND a.private=0 AND a.color IN ('yellow','green') "
        "AND a.quote!='' ORDER BY a.page, a.id LIMIT 12", (pid,))]
    return p


def _ai_points(con, p: dict) -> dict:
    ds = {d["key"]: d for d in params.defs(con)}
    pr = "；".join(f"{ds.get(k, {}).get('name', k)} = {v['value']} {ds.get(k, {}).get('unit', '')}" for k, v in p["params"].items())
    user = (f"論文：{p['title']}（{p['venue']} {p['year'] or ''}）\n重點欄：{json.dumps(p['keyinfo'], ensure_ascii=False)}\n參數：{pr}\n"
            f"標註（重點）：{' / '.join((h['quote'] + (' — ' + h['body'] if h['body'] else ''))[:200] for h in p['highlights'])}\n"
            f"疑問：{' / '.join((q['quote'] + ' ' + q['body'])[:200] for q in p['questions'])}\n摘要：{p['abstract'][:1500]}")
    system = ("你幫實驗室成員準備組會報告投影片。用繁體中文（台灣用語），專有名詞保留英文。只輸出 JSON："
              "{\"one_line\": \"一句話說這篇做了什麼（40 字內）\", \"background\": [...], \"setup\": [...], \"results\": [...], "
              "\"relevance\": [...], \"questions\": [...]}，每個陣列 2–5 點，每點 45 字內，具體、有數值就寫數值。資料沒有的不要編。")
    try:
        out = ai._json_from(ai.chat(con, system, user, 1800))
        return out if isinstance(out, dict) else {}
    except Exception:  # noqa: BLE001 - AI 失敗就退回重點欄
        return {}


def _points(p: dict, use_ai: bool, con) -> dict:
    ki = p["keyinfo"]
    pts = _ai_points(con, p) if use_ai else {}
    def get(k, fallback):
        v = pts.get(k)
        return [str(x)[:90] for x in v][:5] if isinstance(v, list) and v else fallback
    return {
        "one_line": str(pts.get("one_line") or "")[:80] or (_bullets(ki.get("重點", ""), 1, 80) or [""])[0],
        "background": get("background", _bullets(ki.get("重點", ""), 4)),
        "setup": get("setup", _bullets(ki.get("架設／平台", ""), 4)),
        "results": get("results", _bullets(ki.get("可萃取特徵／觀測量", ""), 4)),
        "relevance": get("relevance", _bullets(ki.get("與我們實驗的關係", ""), 4)),
        "questions": get("questions", [((q["body"] or q["quote"])[:90]) for q in p["questions"]][:5]),
    }


# ------------------------------------------------------------------ 各種頁面
def _slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, NAVY if dark else WHITE)
    return s


def s_cover(prs, title, sub, lines):
    s = _slide(prs, dark=True)
    _text(s, M, Inches(2.2), Inches(9.5), Inches(1.4), [(title, 40, WHITE, True)], anchor=MSO_ANCHOR.BOTTOM)
    _text(s, M, Inches(3.75), Inches(9.5), Inches(0.6), [(sub, 18, SOFT)])
    if lines:
        _text(s, M, Inches(4.7), Inches(11), Inches(2.2), [(x, 14, SOFT) for x in lines[:8]])
    return s


def s_agenda(prs, items):
    s = _slide(prs)
    _title(s, "今天的內容")
    y = Inches(1.7)
    for i, (who, title, cite) in enumerate(items[:8]):
        c = _box(s, M, y, Inches(0.55), Inches(0.55), ACCENT, radius=True)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = Emu(0)
        r = tf.paragraphs[0].add_run(); r.text = str(i + 1); _font(r, 16, WHITE, True)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(s, M + Inches(0.8), y - Inches(0.02), Inches(10.8), Inches(0.62),
              [(title, 17, INK, True), (f"{who or '報告人未定'} · {cite}", 12, MUTED)], autofit=True)
        y += Inches(0.68)
    return s


def s_paper_title(prs, con, p, presenter, one_line):
    s = _slide(prs, dark=True)
    img = _page_image(con, p["id"])
    tw = Inches(7.6) if img else W - 2 * M
    au = "、".join(p["authors"][:6]) + (" 等" if len(p["authors"]) > 6 else "")
    ids = "  ".join(x for x in (f"DOI {p['doi']}" if p["doi"] else "", f"arXiv:{p['arxiv']}" if p["arxiv"] else "") if x)
    _text(s, M, Inches(0.9), tw, Inches(2.6), [(p["title"], 30, WHITE, True)], anchor=MSO_ANCHOR.BOTTOM)
    _text(s, M, Inches(3.7), tw, Inches(1.6), [(au, 14, SOFT), (f"{p['venue'] or ''} {p['year'] or ''}".strip(), 14, SOFT), (ids, 11, SOFT)])
    if one_line:
        _text(s, M, Inches(5.4), tw, Inches(0.9), [(one_line, 16, WHITE, False)])
    _text(s, M, H - Inches(0.9), tw, Inches(0.4), [(f"報告人：{presenter or '未定'}　·　{p['citekey']}", 12, SOFT)], autofit=False)
    if img:
        x, y, w, h = Inches(8.7), Inches(0.6), Inches(4.0), Inches(6.3)
        pic = _picture_fit(s, img, x, y, w, h)
        pic.line.color.rgb = SOFT
    return s


def s_points(prs, p, pts, fig):
    """重點頁：左邊條列（背景＋架設），右邊一張圖卡。"""
    s = _slide(prs)
    _title(s, "這篇在做什麼", p["citekey"])
    lw = Inches(6.6) if fig else W - 2 * M
    paras = []
    if pts["background"]:
        paras += [("背景與結果", 15, ACCENT, True)] + [(x, 15, INK, False, PP_ALIGN.LEFT, True) for x in pts["background"]]
    if pts["setup"]:
        paras += [("架設／平台", 15, ACCENT, True)] + [(x, 15, INK, False, PP_ALIGN.LEFT, True) for x in pts["setup"]]
    if not paras:
        paras = [("（重點欄還沒填，可以先用「AI 預填」或手動填寫）", 15, MUTED)]
    _text(s, M, Inches(1.65), lw, Inches(5.2), paras)
    if fig:
        _figure_block(s, fig, p, Inches(7.5), Inches(1.6), Inches(5.2), Inches(5.3))
    _notes(s, "\n".join(f"{k}：{v}" for k, v in p["keyinfo"].items()))
    return s


def _figure_block(s, fig, p, x, y, w, h):
    path = figures.file_path(fig)
    if not path.exists():
        return
    _box(s, x, y, w, h, PALE, radius=True)
    cap_h = Inches(0.9) if fig["caption"] else Inches(0.35)
    pad = Inches(0.15)
    _picture_fit(s, path, x + pad, y + pad, w - 2 * pad, h - cap_h - 2 * pad)
    cap = (fig["caption"][:160] + ("…" if len(fig["caption"]) > 160 else "")) if fig["caption"] else ""
    _text(s, x + pad, y + h - cap_h - Inches(0.05), w - 2 * pad, cap_h,
          ([(cap, 10, INK)] if cap else []) + [(f"{p['citekey']} p.{fig['page']}", 9, MUTED)])


def s_params(prs, con, p):
    ds = {d["key"]: d for d in params.defs(con)}
    rows = [(ds[k], v) for k, v in sorted(p["params"].items(), key=lambda kv: ds.get(kv[0], {}).get("sort", 999)) if k in ds][:11]
    if not rows:
        return None
    s = _slide(prs)
    _title(s, "關鍵參數", p["citekey"] + "（頻率與速率為 ω/2π，線寬為 HWHM）")
    n = len(rows) + 1
    tbl = s.shapes.add_table(n, 4, M, Inches(1.7), W - 2 * M, Inches(0.42) * n).table
    widths = [Inches(3.6), Inches(1.6), Inches(3.4), W - 2 * M - Inches(8.6)]
    for i, wd in enumerate(widths):
        tbl.columns[i].width = wd
    for j, t in enumerate(["參數", "符號", "數值", "條件／備註"]):
        c = tbl.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        r = c.text_frame.paragraphs[0].add_run(); r.text = t; _font(r, 12, WHITE, True)
    for i, (d, v) in enumerate(rows, 1):
        vals = [d["name"], d["symbol"], f"{v['value']} {d['unit']}".strip(), (v["note"] or v["raw"] or "")[:70]]
        for j, t in enumerate(vals):
            c = tbl.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = PALE if i % 2 else WHITE
            r = c.text_frame.paragraphs[0].add_run(); r.text = t; _font(r, 12, INK, j == 2)
    return s


def s_figure(prs, p, fig):
    path = figures.file_path(fig)
    if not path.exists():
        return None
    s = _slide(prs)
    kind = figures.KINDS.get(fig["kind"], "圖")
    m = figures._CAP.match(fig["caption"] or "")
    label = f"{m.group(1).capitalize().rstrip('.')}. {m.group(2)}" if m else kind
    _title(s, fig["note"][:40] if fig["note"] and len(fig["note"]) <= 40 else label, f"{p['citekey']} · 第 {fig['page']} 頁" + (f" · {label}" if m and fig["note"] else ""))
    cap_h = Inches(1.0) if fig["caption"] or fig["note"] else Inches(0.2)
    _picture_fit(s, path, M, Inches(1.6), W - 2 * M, H - Inches(1.6) - cap_h - Inches(0.35))
    if fig["caption"] or fig["note"]:
        cap = fig["caption"][:300] + ("…" if len(fig["caption"]) > 300 else "")
        _text(s, M, H - cap_h - Inches(0.25), W - 2 * M, cap_h,
              ([(cap, 11, MUTED)] if cap else []) + ([(f"筆記：{fig['note'][:160]}", 12, INK, True)] if fig["note"] and len(fig["note"]) > 40 else []))
    if fig["latex"]:
        _notes(s, "LaTeX：" + fig["latex"])
    return s


def s_two_cards(prs, p, pts):
    if not (pts["results"] or pts["relevance"]):
        return None
    s = _slide(prs)
    _title(s, "看得到什麼、跟我們有什麼關係", p["citekey"])
    cw = (W - 2 * M - Inches(0.4)) // 2
    for i, (head, items) in enumerate((("可萃取特徵／觀測量", pts["results"]), ("與我們實驗的關係", pts["relevance"]))):
        x = M + i * (cw + Inches(0.4))
        _box(s, x, Inches(1.7), cw, Inches(5.1), PALE, radius=True)
        _text(s, x + Inches(0.35), Inches(2.0), cw - Inches(0.7), Inches(4.6),
              [(head, 18, NAVY, True)] + [(t, 15, INK, False, PP_ALIGN.LEFT, True) for t in (items or ["（還沒填）"])])
    return s


def s_questions(prs, p, pts):
    qs = pts["questions"]
    if not qs:
        return None
    s = _slide(prs)
    _title(s, "疑問與討論", p["citekey"] + "（來自大家標成「疑問」的標註）")
    _text(s, M, Inches(1.7), W - 2 * M, Inches(5.2), [(q, 17, INK, False, PP_ALIGN.LEFT, True) for q in qs])
    _notes(s, "\n\n".join(f"p.{q['page'] or '-'} {q['who'] or ''}：「{q['quote']}」 {q['body']}" for q in p["questions"]))
    return s


def s_end(prs, text="討論"):
    s = _slide(prs, dark=True)
    _text(s, M, Inches(2.8), W - 2 * M, Inches(1.4), [(text, 44, WHITE, True, PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    return s


def paper_slides(prs, con, pid, presenter, uid, use_ai, fig_ids=None, progress=lambda m: None):
    p = _paper_data(con, pid, uid, fig_ids)
    progress(f"整理「{p['citekey']}」" + ("（AI 精簡中）" if use_ai else ""))
    pts = _points(p, use_ai, con)
    s_paper_title(prs, con, p, presenter, pts["one_line"])
    s_points(prs, p, pts, p["figures"][0] if p["figures"] else None)
    s_params(prs, con, p)
    for f in p["figures"][1:]:
        s_figure(prs, p, f)
    s_two_cards(prs, p, pts)
    s_questions(prs, p, pts)
    return p


def _new():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def _save(prs, name: str) -> str:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[\\/:*?\"<>|\s]+", "_", name)[:80]
    fn = f"{safe}-{datetime.now():%Y%m%d-%H%M%S}.pptx"
    prs.save(OUT_DIR / fn)
    # 只留最近 40 份
    olds = sorted(OUT_DIR.glob("*.pptx"), key=lambda f: f.stat().st_mtime)
    for f in olds[:-40]:
        f.unlink(missing_ok=True)
    return fn


def meeting_deck(con, mid: int, uid: int, use_ai: bool, figs: dict | None, progress=lambda m: None) -> str:
    m = con.execute("SELECT * FROM meetings WHERE id=?", (mid,)).fetchone()
    if m is None:
        raise ValueError("找不到這場組會")
    items = [dict(r) for r in con.execute(
        "SELECT mi.*, u.display_name AS presenter, p.title, p.citekey FROM meeting_items mi LEFT JOIN users u ON u.id=mi.presenter_id "
        "LEFT JOIN papers p ON p.id=mi.paper_id WHERE mi.meeting_id=? ORDER BY mi.sort, mi.id", (mid,))]
    if not items:
        raise ValueError("這場組會還沒有排論文")
    prs = _new()
    presenters = list(dict.fromkeys(i["presenter"] for i in items if i["presenter"]))
    s_cover(prs, m["title"], f"{m['date']}　·　{'、'.join(presenters) or ''}", [])
    s_agenda(prs, [(i["presenter"], i["title"] or i["note"] or "自由主題", i["citekey"] or "") for i in items])
    for it in items:
        if it["paper_id"]:
            f = (figs or {}).get(str(it["paper_id"]))
            paper_slides(prs, con, it["paper_id"], it["presenter"], uid, use_ai, f, progress)
        else:
            s = _slide(prs, dark=True)
            _text(s, M, Inches(2.6), W - 2 * M, Inches(1.6), [(it["note"] or "自由主題", 34, WHITE, True), (f"報告人：{it['presenter'] or '未定'}", 16, SOFT)])
    s_end(prs)
    return _save(prs, f"{m['date']}-{m['title']}")


def paper_deck(con, pid: int, uid: int, use_ai: bool, fig_ids=None, presenter="", progress=lambda m: None) -> str:
    prs = _new()
    p = paper_slides(prs, con, pid, presenter, uid, use_ai, fig_ids, progress)
    s_end(prs)
    return _save(prs, p["citekey"])


def figures_deck(con, ids: list[int], title="圖表剪貼簿") -> str:
    prs = _new()
    rows = [dict(r) for r in con.execute(f"SELECT f.*, p.citekey, p.title AS ptitle FROM figures f JOIN papers p ON p.id=f.paper_id "
                                         f"WHERE f.id IN ({','.join('?' * len(ids)) or 'NULL'})", ids)]
    order = {i: n for n, i in enumerate(ids)}
    rows.sort(key=lambda r: order.get(r["id"], 0))
    if not rows:
        raise ValueError("沒有選到圖卡")
    s_cover(prs, title, f"{datetime.now():%Y-%m-%d}　·　{len(rows)} 張", list(dict.fromkeys(f"{r['citekey']}：{r['ptitle'][:70]}" for r in rows))[:8])
    for r in rows:
        s_figure(prs, {"citekey": r["citekey"]}, r)
    return _save(prs, title)


@handler("slides")
def _job(con, arg, progress):
    a = json.loads(arg or "{}")
    if a.get("meeting"):
        fn = meeting_deck(con, int(a["meeting"]), a.get("user"), bool(a.get("ai")), a.get("figs"), progress)
    elif a.get("paper"):
        fn = paper_deck(con, int(a["paper"]), a.get("user"), bool(a.get("ai")), a.get("fig_ids"), a.get("presenter", ""), progress)
    else:
        fn = figures_deck(con, [int(x) for x in a.get("figures", [])], a.get("title") or "圖表剪貼簿")
    return f"FILE:{fn}"


def export_path(name: str) -> Path:
    p = (OUT_DIR / Path(name).name).resolve()
    if OUT_DIR.resolve() not in p.parents or not p.exists():
        raise FileNotFoundError(name)
    return p


_ = (io, db)
