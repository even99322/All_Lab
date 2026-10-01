"""PowerPoint export: every part is a separate, editable shape with PowerPoint's own 3-D.

Board layers are flat freeform outlines (screw holes and vias cut out) that get
PowerPoint "3-D Rotation" (orthographic camera tilted back by 90 deg - elevation)
and "3-D Format" (depth = layer thickness, material, lighting). The azimuth is
applied to the outlines themselves, so every layer uses the same tilt and the
parts line up as one assembly: each shape is placed where the orthographic
camera puts its centre (the same projection as the editor).

YIG  = circle with round top and bottom bevels (a real 3-D sphere) + white spin shapes.
Coil = one donut per turn with round bevels (a tube ring), tilted like the board,
       plus an up arrow.
Wave = an open freeform line (dashed), freely resizable in PowerPoint.
Each YIG / coil / layer is a group named after it; everything stays editable
(Format Shape > 3-D Format / 3-D Rotation, colours, lighting, material).
"""

from __future__ import annotations

import math
from pathlib import Path

from lxml import etree
from PySide6.QtCore import QRectF
from shapely.geometry.polygon import orient

from app.figure_builder.dxf_import import polygons
from app.figure_builder.render import coil_turn_heights, make_frame, spin_positions, wave_points
from app.figure_builder.scene import Projection, Scene

SLIDE_W = 12192000          # 13.333 in (16:9), EMU
SLIDE_H = 6858000
EMU_PER_PT = 12700
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _q(tag: str) -> str:
    return f"{{{A}}}{tag}"


def _hex(color: str) -> str:
    return color.lstrip("#").upper()[:6]


def _darker(color: str, factor: float = 0.6) -> str:
    value = _hex(color)
    r, g, b = (int(value[i:i + 2], 16) for i in (0, 2, 4))
    return f"{int(r * factor):02X}{int(g * factor):02X}{int(b * factor):02X}"


def _sp_pr(shape):
    return shape._element.spPr


def _insert_3d(shape, *, tilt_deg: float, depth_emu: int = 0, material: str = "plastic", bevel_emu: int = 0,
               bevel_h_emu: int | None = None, extrusion_color: str | None = None, rotate: bool = True) -> None:
    """Add <a:scene3d> and <a:sp3d> to the shape (in schema order, after fill / line / effects)."""
    sp_pr = _sp_pr(shape)
    for old in sp_pr.findall(_q("scene3d")) + sp_pr.findall(_q("sp3d")):
        sp_pr.remove(old)
    if sp_pr.find(_q("effectLst")) is None:
        etree.SubElement(sp_pr, _q("effectLst"))       # no theme shadow; users can add effects in PowerPoint
    scene = etree.SubElement(sp_pr, _q("scene3d"))
    camera = etree.SubElement(scene, _q("camera"), prst="orthographicFront")
    if rotate and tilt_deg:
        # PowerPoint "Y rotation": 360 - tilt lays the shape back like a floor (top edge away from the viewer).
        lat = int(round(((360.0 - tilt_deg) % 360.0) * 60000))
        etree.SubElement(camera, _q("rot"), lat=str(lat), lon="0", rev="0")
    light = etree.SubElement(scene, _q("lightRig"), rig="threePt", dir="t")
    etree.SubElement(light, _q("rot"), lat="0", lon="0", rev="1200000")
    sp3d = etree.SubElement(sp_pr, _q("sp3d"), prstMaterial=material)
    if depth_emu > 0:
        sp3d.set("extrusionH", str(int(depth_emu)))
    if bevel_emu > 0:
        h = str(int(bevel_h_emu if bevel_h_emu is not None else bevel_emu))
        etree.SubElement(sp3d, _q("bevelT"), w=str(int(bevel_emu)), h=h, prst="circle")
        etree.SubElement(sp3d, _q("bevelB"), w=str(int(bevel_emu)), h=h, prst="circle")
    if extrusion_color:
        clr = etree.SubElement(sp3d, _q("extrusionClr"))
        etree.SubElement(clr, _q("srgbClr"), val=extrusion_color)


def _style(shape, fill: str | None, line: str | None = None, line_pt: float = 0.0) -> None:
    if fill is None:
        shape.fill.background()
    else:
        from pptx.dml.color import RGBColor

        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(_hex(fill))
    if line is None:
        shape.line.fill.background()
    else:
        from pptx.dml.color import RGBColor
        from pptx.util import Pt

        shape.line.color.rgb = RGBColor.from_string(_hex(line))
        shape.line.width = Pt(line_pt)


def _freeform(shapes, rings: list[list[tuple[float, float]]], close: bool = True):
    """Freeform from rings of slide EMU points (one path, one sub-path per ring)."""
    first = rings[0]
    builder = shapes.build_freeform(int(round(first[0][0])), int(round(first[0][1])), scale=1.0)
    builder.add_line_segments([(int(round(x)), int(round(y))) for x, y in first[1:]], close=close)
    for ring in rings[1:]:
        builder.move_to(int(round(ring[0][0])), int(round(ring[0][1])))
        builder.add_line_segments([(int(round(x)), int(round(y))) for x, y in ring[1:]], close=close)
    return builder.convert_to_shape()


def _placed_rings(polygon, proj: Projection, frame, z: float):
    """Outline rings in the board plane (azimuth applied), positioned so the bounding-box centre
    lands where the orthographic camera shows it; PowerPoint's tilt does the foreshortening."""
    polygon = orient(polygon, 1.0)
    rings = [[proj.rotate(x, y) for x, y in ring.coords[:-1]] for ring in [polygon.exterior, *polygon.interiors]]
    xs = [p[0] for ring in rings for p in ring]
    ys = [p[1] for ring in rings for p in ring]
    cx, cy = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
    center = frame.point(cx, cy * proj.sp + z * proj.cp)
    s = frame.scale
    return [[(center.x() + (x - cx) * s, center.y() - (y - cy) * s) for x, y in ring] for ring in rings]


def export_pptx(scene: Scene, path: str | Path) -> Path:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu

    scene.update_levels()
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])        # blank
    proj = Projection(scene.camera)
    frame = make_frame(scene, QRectF(0, 0, SLIDE_W, SLIDE_H), proj)
    tilt = 90.0 - max(1.0, min(90.0, scene.camera.elevation))
    s = frame.scale

    # coils first: they sit under the board
    for coil in scene.coils:
        group = slide.shapes.add_group_shape()
        group.name = f"Coil {coil.id}"
        heights = coil_turn_heights(scene, coil)
        size = 2 * coil.radius * s
        for index, z in enumerate(reversed(heights)):
            center = frame.point(*proj.project(coil.x, coil.y, z))
            ring = group.shapes.add_shape(MSO_SHAPE.DONUT, Emu(int(center.x() - size / 2)),
                                          Emu(int(center.y() - size / 2)), Emu(int(size)), Emu(int(size)))
            ring.name = f"Coil {coil.id} turn {index + 1}"
            ring.adjustments[0] = min(0.5, coil.wire / (2 * coil.radius))
            _style(ring, coil.color)
            _insert_3d(ring, tilt_deg=tilt, material="metal", bevel_emu=coil.wire / 2 * s,
                       extrusion_color=_darker(coil.color))
        if coil.arrow:
            bottom = frame.point(*proj.project(coil.x, coil.y, heights[-1]))
            top = frame.point(*proj.project(coil.x, coil.y, scene.bottom() - coil.gap * 0.2))
            width = coil.radius * s * 0.55
            arrow = group.shapes.add_shape(MSO_SHAPE.UP_ARROW, Emu(int(top.x() - width / 2)), Emu(int(top.y())),
                                           Emu(int(width)), Emu(int(max(1, bottom.y() - top.y()))))
            arrow.name = f"Coil {coil.id} field arrow"
            _style(arrow, coil.arrow_color, _darker(coil.arrow_color), 0.75)
            _insert_3d(arrow, tilt_deg=0, material="plastic", bevel_emu=width * 0.08, rotate=False)

    # board layers, bottom up
    order = sorted((l for l in scene.layers if l.visible), key=lambda l: (l.z1, l.placement == "through"))
    for layer in order:
        parts = polygons(scene.solid(layer))
        if not parts:
            continue
        group = slide.shapes.add_group_shape()
        group.name = layer.name
        for index, polygon in enumerate(parts):
            shape = _freeform(group.shapes, _placed_rings(polygon, proj, frame, layer.z1))
            shape.name = f"{layer.name} {index + 1}" if len(parts) > 1 else layer.name
            _style(shape, layer.color)
            _insert_3d(shape, tilt_deg=tilt, depth_emu=layer.thickness * s, material=layer.material,
                       extrusion_color=_darker(layer.color, 0.75))

    # YIG spheres
    for yig in sorted(scene.yigs, key=lambda y: proj.depth(y.x, y.y, 0)):
        group = slide.shapes.add_group_shape()
        group.name = f"YIG {yig.id}"
        surface = scene.surface_height(yig.x, yig.y)
        center = frame.point(*proj.project(yig.x, yig.y, surface + yig.lift + yig.radius))
        d = 2 * yig.radius * s
        sphere = group.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(center.x() - d / 2)), Emu(int(center.y() - d / 2)),
                                        Emu(int(d)), Emu(int(d)))
        sphere.name = f"YIG {yig.id} sphere"
        _style(sphere, yig.color)
        _insert_3d(sphere, tilt_deg=0, material="plastic", bevel_emu=d / 2, rotate=False)
        glyph = yig.radius * s * 0.3
        for n, (u, v) in enumerate(spin_positions(yig.spins)):
            gx, gy = center.x() + u * yig.radius * s, center.y() + v * yig.radius * s
            head = [(gx, gy - glyph * 0.55), (gx + glyph * 0.2, gy - glyph * 0.22), (gx - glyph * 0.2, gy - glyph * 0.22)]
            arrow = _freeform(group.shapes, [head])
            _style(arrow, yig.spin_color)
            arrow.name = f"YIG {yig.id} spin {n + 1} head"
            stem = _freeform(group.shapes, [[(gx, gy + glyph * 0.5), (gx, gy - glyph * 0.3)]], close=False)
            _style(stem, None, yig.spin_color, max(0.5, glyph / EMU_PER_PT * 0.12))
            stem.name = f"YIG {yig.id} spin {n + 1} stem"
            ring = group.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(gx - glyph * 0.28)), Emu(int(gy + glyph * 0.12)),
                                          Emu(int(glyph * 0.56)), Emu(int(glyph * 0.2)))
            _style(ring, None, yig.spin_color, max(0.5, glyph / EMU_PER_PT * 0.1))
            ring.name = f"YIG {yig.id} spin {n + 1} cone"

    # sine waves (free size)
    for wave in scene.waves:
        points = [(p.x(), p.y()) for p in wave_points(wave, frame, 160)]
        shape = _freeform(slide.shapes, [points], close=False)
        shape.name = f"Sine wave {wave.id}"
        _style(shape, None, wave.color, wave.width)
        if wave.dashed:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE

            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    path = Path(path)
    presentation.save(str(path))
    return path


def slide_shape_summary(path: str | Path) -> list[tuple[str, str]]:
    """(name, kind) of every shape, groups expanded (used by tests)."""
    from pptx import Presentation

    rows = []

    def walk(shapes):
        for shape in shapes:
            rows.append((shape.name, shape.shape_type and str(shape.shape_type)))
            if shape.shape_type is not None and "GROUP" in str(shape.shape_type):
                walk(shape.shapes)

    walk(Presentation(str(path)).slides[0].shapes)
    return rows
